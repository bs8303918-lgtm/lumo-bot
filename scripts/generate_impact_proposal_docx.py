"""Generate Impact Admission commercial proposal as .docx."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "IMPACT_ADMISSION_proposal.docx"


def set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Коммерческое предложение")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x07, 0x07, 0x0D)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Lumo × Impact Admission")
    r2.font.size = Pt(16)
    r2.font.color.rgb = RGBColor(0x22, 0xD3, 0xEE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m = meta.add_run("Июнь 2026")
    m.font.size = Pt(10)
    m.font.color.rgb = RGBColor(0x71, 0x71, 0x7A)

    doc.add_paragraph()


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullet(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
    doc.add_paragraph()


def add_quote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = RGBColor(0x3F, 0x3F, 0x46)


def build() -> Path:
    doc = Document()
    set_doc_defaults(doc)
    add_title_block(doc)

    # 1. Summary
    heading(doc, "1. Кратко")
    doc.add_paragraph(
        "Lumo — Telegram-бот с ИИ, который 24/7 мониторит каналы с грантами, "
        "стипендиями, олимпиадами, летними школами, конкурсами и стажировками "
        "и присылает ученику только то, что подходит под его профиль."
    )
    doc.add_paragraph(
        "Impact Admission помогает школьникам выигрывать возможности и поступать "
        "зарубеж — стратегия, сопровождение, эссе, выбор программ."
    )
    add_quote(
        doc,
        "Impact занимается стратегией поступления. "
        "Lumo занимается поиском возможностей и дедлайнов.",
    )

    # 2. Problem
    heading(doc, "2. Проблема")
    add_table(
        doc,
        ["Проблема", "Последствие"],
        [
            ["Сотни Telegram-каналов и сайтов с программами", "Консультанты тратят часы на ручной поиск"],
            ["Дедлайны разбросаны по постам и языкам", "Ученики пропускают гранты, олимпиады, летние школы"],
            ["У каждого ученика свой профиль", "Общая рассылка «топ-10 программ» не работает"],
            ["Конкуренты добавляют технологичность", "Сложнее выделиться только консультациями"],
        ],
    )
    doc.add_paragraph(
        "Пропущенная летняя школа или олимпиада — слабее портфолио при поступлении. "
        "Impact не может физически следить за всеми каналами для каждого клиента."
    )

    # 3. Solution
    heading(doc, "3. Решение — Lumo")
    heading(doc, "Что делает Lumo", level=2)
    for item in [
        "Мониторинг — автоматически отслеживает Telegram-каналы с возможностями.",
        "ИИ-классификация — извлекает название, тип, дедлайн, требования и ссылку; фильтрует рекламу.",
        "Персонализация — ученик описывает себя один раз, дальше бот присылает только релевантное.",
        "Доставка в Telegram — подборки приходят туда, где ученик уже каждый день.",
    ]:
        add_bullet(doc, item)

    heading(doc, "Что Lumo не делает", level=2)
    for item in [
        "Не пишет мотивационные письма и эссе",
        "Не выбирает вузы и не строит стратегию поступления",
        "Не готовит к SAT / IELTS",
        "Не заменяет консультанта Impact — дополняет его работу",
    ]:
        add_bullet(doc, item)

    # 4. Roles
    heading(doc, "4. Разделение ролей")
    add_table(
        doc,
        ["Impact Admission", "Lumo"],
        [
            ["Стратегия поступления", "Поиск возможностей"],
            ["Эссе и документы", "Мониторинг каналов"],
            ["Выбор программ", "Персональные подборки"],
            ["Сопровождение ученика", "Контроль дедлайнов"],
        ],
    )

    # 5. Packages - MAIN
    heading(doc, "5. Коммерческие пакеты")
    doc.add_paragraph(
        "Предлагаем два формата. Рекомендуем начать с Пакета 1 — быстрый пилот без интеграции."
    )

    heading(doc, "Пакет 1 — «30 подписок в Lumo» (пилот, рекомендуем)", level=2)
    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ["Цена", "29 890 ₸ — фикс за пакет"],
            ["Что входит", "30 активаций подписки в Telegram-боте Lumo на 1 месяц"],
            ["Кому", "Impact сам решает: ученикам, консультантам или смешанно"],
            ["Как работает", "Impact получает партнёрскую ссылку → раздаёт людям → Lumo активирует подписку"],
            ["Что получает пользователь", "Безлимитный AI-поиск, персональные подборки, уведомления о возможностях и дедлайнах"],
            ["Срок", "30 календарных дней с момента активации пакета"],
        ],
    )
    doc.add_paragraph(
        "Зачем Impact: быстро проверить ценность без интеграции в платформу. "
        "Можно оформить как бонус клиентам или внутреннюю акцию для команды."
    )

    heading(doc, "Пакет 2 — «База конкурсов через API»", level=2)
    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ["Цена", "49 850 ₸ / месяц"],
            ["Что входит", "Доступ к API с актуальной базой конкурсов, грантов и программ"],
            ["Как работает", "Lumo мониторит каналы → ИИ классифицирует посты → база пополняется автоматически → Impact забирает данные через API"],
            ["Обновления", "Без ручной работы: новые конкурсы и дедлайны появляются в API сами"],
            ["Техобслуживание", "Мониторинг каналов, актуализация базы, поддержка API — на стороне Lumo"],
            ["Срок подключения", "3–7 дней (выдача API-ключа + документация; настройка sync — на стороне Impact)"],
        ],
    )
    doc.add_paragraph("Что делает Lumo: бот работает, база растёт, API отдаёт JSON.")
    doc.add_paragraph(
        "Что делает Impact: подключает API к своей платформе (cron / sync раз в час) "
        "и показывает конкурсы ученикам."
    )

    heading(doc, "Сравнение пакетов", level=2)
    add_table(
        doc,
        ["", "Пакет 1", "Пакет 2"],
        [
            ["Цена", "29 890 ₸ (разово)", "49 850 ₸ / мес"],
            ["Срок запуска", "1–3 дня", "3–7 дней"],
            ["Для кого", "Ученики или консультанты", "Платформа Impact (все ученики)"],
            ["Где живёт", "Telegram-бот", "Сайт / ЛК Impact (данные из API)"],
            ["Кто настраивает sync", "—", "Impact (1 раз), Lumo даёт API"],
            ["Лучше для", "Быстрый пилот", "Долгосрочный продукт"],
        ],
    )

    # 6. Pilot
    heading(doc, "6. Что измеряем в пилоте (Пакет 1)")
    for item in [
        "Сколько из 30 человек активировали бота",
        "Сколько персональных подборок получено",
        "Обратная связь от 3–5 пользователей и 1–2 сотрудников Impact",
    ]:
        add_bullet(doc, item)
    doc.add_paragraph(
        "По итогам месяца — решение: продлить пакет, перейти на Пакет 2 или масштабировать количество мест."
    )

    # 7. Technical
    heading(doc, "7. Техническая схема (Пакет 2)")
    scheme = (
        "Telegram-каналы\n"
        "      ↓\n"
        "Lumo бот (мониторинг + ИИ-классификация)\n"
        "      ↓\n"
        "База конкурсов (автообновление)\n"
        "      ↓\n"
        "Partner API → GET /catalog/opportunities\n"
        "      ↓\n"
        "Платформа Impact (sync раз в N часов)\n"
        "      ↓\n"
        "Ученики видят актуальные конкурсы"
    )
    p = doc.add_paragraph()
    r = p.add_run(scheme)
    r.font.name = "Consolas"
    r.font.size = Pt(10)

    doc.add_paragraph(
        "Lumo не встраивается в код Impact — отдаёт готовый JSON по API. "
        "Impact один раз настраивает автозагрузку на своей стороне."
    )

    add_table(
        doc,
        ["Endpoint", "Назначение"],
        [
            ["GET /api/partner/v1/catalog/opportunities", "Список конкурсов (title, type, deadline, description, url)"],
            ["GET /api/partner/v1/health", "Проверка доступности"],
            ["Authorization: Bearer <API_KEY>", "Доступ только для Impact"],
        ],
    )

    # 8. Student text
    heading(doc, "8. Готовый текст для рассылки ученикам (Пакет 1)")
    text = (
        "Привет! Мы подключили для вас Lumo — AI-помощника, который ищет "
        "гранты, олимпиады, летние школы и конкурсы под ваш профиль.\n\n"
        "Как начать:\n"
        "1. Откройте бота: [ссылка Impact]\n"
        "2. Нажмите Start\n"
        "3. Отправьте команду /set_interest и напишите о себе: "
        "возраст, класс, интересы, куда хотите поступать\n\n"
        "Бот будет присылать подборки в Telegram. "
        "Ваш консультант Impact по-прежнему ведёт стратегию поступления — "
        "Lumo помогает не пропустить возможности и дедлайны."
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    r = p.add_run(text)
    r.italic = True

    # 9. Next steps
    heading(doc, "9. Следующие шаги")
    add_table(
        doc,
        ["#", "Действие", "Срок"],
        [
            ["1", "Демо Lumo (15–30 мин)", "по согласованию"],
            ["2", "Согласование Пакета 1 (30 подписок)", "1 встреча"],
            ["3", "Impact рассылает ссылку ученикам / команде", "день 1"],
            ["4", "Промежуточный созвон (опционально)", "день 15"],
            ["5", "Отчёт по пилоту + решение о Пакете 2", "день 30"],
            ["6", "Коммерческий договор", "после пилота"],
        ],
    )

    # 10. Contacts
    heading(doc, "10. Контакты")
    add_bullet(doc, " Telegram-бот: t.me/LumoAI1bot", bold_prefix="")
    p = doc.add_paragraph(style="List Bullet")
    p.add_run("Telegram-бот: ").bold = False
    link = p.add_run("t.me/LumoAI1bot")
    link.font.color.rgb = RGBColor(0x22, 0xD3, 0xEE)

    p2 = doc.add_paragraph(style="List Bullet")
    p2.add_run("Сообщество: ")
    p2.add_run("t.me/LumoCommunity").font.color.rgb = RGBColor(0x22, 0xD3, 0xEE)

    doc.add_paragraph()
    cp = doc.add_paragraph("Контакт для переговоров: [указать ваш Telegram / email]")
    cp.runs[0].italic = True

    # Appendix
    heading(doc, "Приложение: типы возможностей в базе Lumo")
    for item in [
        "Гранты и стипендии",
        "Олимпиады (в т.ч. с международным признанием)",
        "Летние школы и программы для школьников",
        "Конкурсы, кейс-чемпионаты, эссе-конкурсы",
        "Стажировки",
        "Хакатоны и IT-мероприятия",
        "Курсы и программы развития",
    ]:
        add_bullet(doc, item)
    doc.add_paragraph("ИИ понимает посты на русском, казахском и английском языках.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Saved: {path}")
